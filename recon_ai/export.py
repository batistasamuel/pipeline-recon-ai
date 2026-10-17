import json
from io import BytesIO
from pathlib import Path

from recon_ai import db
from recon_ai.view import run_view


class ExportError(Exception):
    def __init__(self, message, code=2):
        super().__init__(message)
        self.code = code


def _text_blocks(view, output_dir, locale=None):
    from recon_ai.i18n import t

    run = view["run"]
    lines = [
        f"RECON REPORT - {run['target']}",
        f"run_id: {run['id']}",
        f"status: {run['status']}",
        t("export.started", locale, value=run["started_at"]),
        t("export.finished", locale, value=run.get("finished_at") or "—"),
        f"threads: {run['threads']}",
        "",
    ]
    spec = view.get("spec") or {}
    if spec:
        lines.append(t("export.spec", locale))
        lines.append(t("export.scope", locale, value=", ".join(spec.get("in_scope") or []) or "—"))
        lines.append(t("export.out", locale, value=", ".join(spec.get("out_scope") or []) or "—"))
        for name, item in (spec.get("tools") or {}).items():
            enabled = "on" if item.get("enabled", True) else "off"
            lines.append(f"{name}: {enabled}")
        lines.append("")
    lines.append(t("export.flow", locale))
    for item in view.get("flow") or view.get("phases") or []:
        lines.append(f"{item.get('name')}\t{item.get('status')}\t{item.get('detail') or ''}")
    lines.append("")
    resumo = Path(output_dir) / "resumo.txt"
    if resumo.is_file():
        lines.append(resumo.read_text(encoding="utf-8", errors="replace"))
    lines.append(t("export.findings", locale))
    connection_findings = view.get("findings") or []
    if not connection_findings:
        lines.append(t("export.none", locale))
    for item in connection_findings:
        lines.append(f"{item.get('name')}: {item.get('count')}")
    candidates = view.get("candidates") or []
    if candidates:
        lines.append("")
        lines.append(t("export.candidates", locale))
        for row in candidates:
            lines.append(f"{row.get('severity')} {row.get('template_id')} {row.get('matched_at')}")
    checklist = view.get("checklist") or []
    if checklist:
        lines.append("")
        lines.append(t("export.checklist", locale))
        lines.extend(checklist)
    if view.get("classification"):
        lines.append("")
        lines.append(t("export.classification", locale))
        lines.append(json.dumps(view["classification"], ensure_ascii=False))
    if view.get("delta"):
        lines.append("")
        lines.append(t("export.delta", locale))
        for item in view["delta"]:
            lines.append(f"{item.get('change')} {item.get('entity')} {item.get('key')}")
    if view.get("routes"):
        lines.append("")
        lines.append(t("export.routes", locale))
        for item in view["routes"]:
            lines.append(f"{item.get('method') or ''} {item.get('path')}".strip())
    clusters = view.get("clusters") or []
    coverage = [row for row in clusters if row.get("class") == "coverage"]
    if coverage:
        lines.append("")
        lines.append(t("export.clusters", locale))
        for row in coverage:
            lines.append(
                f"{row.get('class')} {row.get('template_id')} n={row.get('count')} {row.get('sample') or ''}"
            )
    return lines


def _load_view(database, results, run_id, locale=None):
    from recon_ai.i18n import display_name, t

    path = Path(database)
    if not path.is_file():
        raise ExportError(t("export.missing_db", locale, path=path), 2)
    connection = db.connect(path)
    view = run_view(connection, run_id, locale)
    if view is None:
        raise ExportError(t("export.missing_run", locale, run_id=run_id), 2)
    run = view["run"]
    if run["status"] == "running":
        raise ExportError(t("export.running", locale), 2)
    output = Path(results) / run["output_dir"]
    findings = connection.execute(
        "SELECT severity, name, matched_at, template_id FROM finding WHERE run_id = ? ORDER BY id",
        (run_id,),
    ).fetchall()
    view["finding_rows"] = []
    for row in findings:
        item = dict(row)
        item["name"] = display_name(item.get("template_id"), item.get("name"), locale)
        view["finding_rows"].append(item)
    view["candidates"] = [
        row for row in view["finding_rows"] if str(row.get("template_id") or "").startswith("corr/")
    ]
    from recon_ai.cluster import read_clusters
    from recon_ai.drafts import read_index

    view["clusters"] = read_clusters(output)
    try:
        view["drafts"] = read_index(output)
    except (OSError, json.JSONDecodeError):
        view["drafts"] = []
    connection.close()
    return view, output


def render_pdf(view, output_dir, locale=None):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    nuclei_rows = [
        row
        for row in view.get("finding_rows") or []
        if not str(row.get("template_id") or "").startswith("corr/")
    ]
    extra = []
    if nuclei_rows:
        extra.append("")
        extra.extend(
            f"{row.get('severity')} {row.get('template_id')} {row.get('matched_at')}"
            for row in nuclei_rows
        )
    buffer = BytesIO()
    page = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    y = height - 48
    for line in _text_blocks(view, output_dir, locale) + extra:
        if y < 48:
            page.showPage()
            y = height - 48
        page.setFont("Times-Roman", 10)
        page.drawString(48, y, line[:110])
        y -= 14
    page.save()
    return buffer.getvalue()


def render_docx(view, output_dir, locale=None):
    from docx import Document

    from recon_ai.i18n import t

    document = Document()
    document.add_heading(f"RECON REPORT - {view['run']['target']}", level=1)
    for line in _text_blocks(view, output_dir, locale):
        document.add_paragraph(line)
    nuclei_rows = [
        row for row in view.get("finding_rows") or []
        if not str(row.get("template_id") or "").startswith("corr/")
    ]
    if nuclei_rows:
        document.add_heading(t("export.findings", locale), level=2)
        for row in nuclei_rows:
            document.add_paragraph(
                f"{row.get('severity')} · {row.get('name')} · {row.get('template_id')} · {row.get('matched_at')}"
            )
    candidates = view.get("candidates") or [
        row for row in view.get("finding_rows") or []
        if str(row.get("template_id") or "").startswith("corr/")
    ]
    if candidates:
        document.add_heading(t("export.candidates", locale), level=2)
        for row in candidates:
            document.add_paragraph(
                f"{row.get('severity')} · {row.get('name')} · {row.get('template_id')} · {row.get('matched_at')}"
            )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _draft_lines(output_dir, cluster_key, locale=None):
    from recon_ai.drafts import DraftError, lookup_key, read_markdown
    from recon_ai.i18n import t

    try:
        key = lookup_key(output_dir, cluster_key)
        text = read_markdown(output_dir, key)
    except DraftError as exc:
        raise ExportError(str(exc), 2) from exc
    if not text.strip():
        raise ExportError(t("export.missing_draft", locale), 2)
    return text.splitlines()


def render_draft_pdf(lines):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = BytesIO()
    page = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    y = height - 48
    for line in lines:
        if y < 48:
            page.showPage()
            y = height - 48
        page.setFont("Times-Roman", 10)
        page.drawString(48, y, line[:110])
        y -= 14
    page.save()
    return buffer.getvalue()


def render_draft_docx(lines, locale=None):
    from docx import Document

    document = Document()
    if lines:
        document.add_heading(lines[0].lstrip("# ").strip() or "Draft", level=1)
        for line in lines[1:]:
            document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_export(view, output_dir, fmt, locale=None, draft=None):
    from recon_ai.i18n import t

    chosen = (fmt or "pdf").lower()
    if draft:
        lines = _draft_lines(output_dir, draft, locale)
        if chosen == "pdf":
            return "application/pdf", render_draft_pdf(lines)
        if chosen in {"docx", "word"}:
            return (
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                render_draft_docx(lines, locale),
            )
        raise ExportError(t("export.bad_format", locale), 2)
    if chosen == "pdf":
        return "application/pdf", render_pdf(view, output_dir, locale)
    if chosen in {"docx", "word"}:
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document", render_docx(
            view, output_dir, locale
        )
    raise ExportError(t("export.bad_format", locale), 2)


def write_export(database, results, run_id, fmt, output=None, locale=None, draft=None):
    view, output_dir = _load_view(database, results, run_id, locale)
    _content_type, body = render_export(view, output_dir, fmt, locale, draft=draft)
    chosen = "docx" if (fmt or "").lower() in {"docx", "word"} else "pdf"
    default_name = f"rascunho.{chosen}" if draft else f"relatorio.{chosen}"
    path = Path(output) if output else Path(results) / view["run"]["output_dir"] / default_name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(body)
    return path
