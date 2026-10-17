#!/usr/bin/env python3
"""Gera o Word consolidado a partir dos Markdown desta pasta."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
ORDER = [
    "README.md",
    "01-overview.md",
    "02-glossary.md",
    "03-history-and-evolution.md",
    "04-scope-and-non-goals.md",
    "architecture/architecture.md",
    "architecture/component-diagram.md",
    "architecture/deployment-diagram.md",
    "architecture/data-flow.md",
    "architecture/execution-flow.md",
    "architecture/sequence/end-to-end.md",
    "architecture/sequence/recon.md",
    "architecture/sequence/js-ai.md",
    "configuration/reference.md",
    "configuration/defaults.md",
    "configuration/parameter-impact.md",
    "configuration/profiles.md",
    "ai/architecture.md",
    "ai/context-window.md",
    "ai/token-budget.md",
    "ai/chunking.md",
    "ai/optimization.md",
    "ai/validation.md",
    "ai/hardware-adaptation.md",
    "operations/installation.md",
    "operations/quick-start.md",
    "operations/usage.md",
    "operations/troubleshooting.md",
    "operations/recovery.md",
    "security/authorized-use.md",
    "security/scope-policy.md",
    "security/data-handling.md",
    "observability/logging.md",
    "observability/run-lifecycle.md",
    "observability/manifest.md",
    "performance/benchmarks.md",
    "performance/optimization.md",
    "development/developer-guide.md",
    "development/testing.md",
    "development/adapters.md",
    "development/extension-guide.md",
    "decisions/ADR-001-docker-contrato.md",
    "decisions/ADR-002-debian-pinado.md",
    "decisions/ADR-003-amd64-arm64.md",
    "decisions/ADR-004-ferramentas-pinadas.md",
    "decisions/ADR-005-cli-typer.md",
    "decisions/ADR-006-sqlite.md",
    "decisions/ADR-007-escopo-obrigatorio.md",
    "decisions/ADR-008-adapters-python.md",
    "decisions/ADR-009-ia-pos-processamento.md",
    "decisions/ADR-010-ollama-e-scan.md",
    "decisions/ADR-011-schema-ia.md",
    "decisions/ADR-012-cloud-opt-in.md",
    "decisions/ADR-013-sem-celery.md",
    "decisions/ADR-014-ui-opcional.md",
    "decisions/ADR-015-coletores-adiados.md",
    "decisions/ADR-016-recon-diferencial.md",
    "decisions/ADR-017-derive-correlate-dorks.md",
    "decisions/ADR-018-reportable-vs-nota.md",
    "appendix/planned-vs-implemented.md",
    "appendix/limitations.md",
    "appendix/documentation-gaps.md",
]


def heading_level(line):
    if line.startswith("# "):
        return 1, line[2:].strip()
    if line.startswith("## "):
        return 2, line[3:].strip()
    if line.startswith("### "):
        return 3, line[4:].strip()
    return None, line


def add_runs(paragraph, text):
    paragraph.add_run(text)


def flush_table(document, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = row[j] if j < len(row) else ""
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True


def parse_table_row(line):
    parts = [p.strip() for p in line.strip().strip("|").split("|")]
    return parts


def is_sep(row):
    return row and all(set(cell.replace(":", "").replace("-", "").strip()) <= set() or set(cell) <= set("-: ") for cell in row)


def convert(path, document, first=False):
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    i = 0
    in_code = False
    code = []
    table = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if in_code:
                block = "\n".join(code)
                paragraph = document.add_paragraph()
                run = paragraph.add_run(block)
                run.font.name = "Courier New"
                run.font.size = Pt(8)
                rpr = run._element.get_or_add_rPr()
                rfonts = rpr.get_or_add_rFonts()
                rfonts.set(qn("w:eastAsia"), "Courier New")
                in_code = False
                code = []
            else:
                in_code = True
                code = []
            i += 1
            continue
        if in_code:
            code.append(line)
            i += 1
            continue
        if line.strip().startswith("|"):
            table.append(parse_table_row(line))
            i += 1
            continue
        if table:
            body = [row for row in table if not is_sep(row)]
            flush_table(document, body)
            table = []
        if not line.strip():
            i += 1
            continue
        level, content = heading_level(line)
        if level:
            if level == 1 and not first and path.name != "README.md":
                document.add_page_break()
            document.add_heading(content, level=min(level, 3))
            first = False
            i += 1
            continue
        if line.startswith("- ") or line.startswith("* "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
            i += 1
            continue
        if len(line) > 2 and line[0].isdigit() and line[1:3] in (". ", ") "):
            document.add_paragraph(line.split(" ", 1)[1], style="List Number")
            i += 1
            continue
        document.add_paragraph(line)
        i += 1
    if table:
        flush_table(document, [row for row in table if not is_sep(row)])
    if in_code and code:
        paragraph = document.add_paragraph()
        run = paragraph.add_run("\n".join(code))
        run.font.name = "Courier New"
        run.font.size = Pt(8)


def main():
    document = Document()
    section = document.sections[0]
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    document.add_heading("Recon AI — documentação consolidada", 0)
    document.add_paragraph(
        "Mesmo conteúdo da pasta documentation/. "
        "Comportamento vigente: código e testes. Intenção: planos na raiz e guia v2. "
        "Gerado a partir dos Markdown, 2026-10-06."
    )
    first = True
    for rel in ORDER:
        path = ROOT / rel
        if not path.is_file():
            raise SystemExit(f"em falta: {rel}")
        convert(path, document, first=first)
        first = False
    out = ROOT / "Recon_AI_Documentacao_Consolidada.docx"
    document.save(out)
    print(out)


if __name__ == "__main__":
    main()
