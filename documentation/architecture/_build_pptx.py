#!/usr/bin/env python3
"""Gera o PPT de arquitetura Compose a partir do estado atual do repositório."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from pathlib import Path

OUT = Path(__file__).resolve().parent / "recon-ai-arquitetura.pptx"

BG = RGBColor(11, 16, 32)
SURFACE = RGBColor(22, 32, 51)
CARD = RGBColor(31, 41, 55)
LINE = RGBColor(55, 65, 81)
CYAN = RGBColor(56, 189, 248)
GREEN = RGBColor(34, 197, 94)
AMBER = RGBColor(245, 158, 11)
RED = RGBColor(239, 68, 68)
PURPLE = RGBColor(167, 139, 250)
WHITE = RGBColor(243, 244, 246)
MUTED = RGBColor(156, 163, 175)
NAVY = RGBColor(15, 23, 42)

W = Inches(13.333)
H = Inches(7.5)


def rgb(shape, fill, line=None, width=Pt(1.25)):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    shape.line.width = width


def set_run(run, text, size, color, bold=False, name="Calibri"):
    run.text = text
    run.font.size = size
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = name


def tf_clear(shape):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    return tf


def add_text(shape, lines, align=PP_ALIGN.LEFT):
    tf = tf_clear(shape)
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        run = p.add_run()
        set_run(run, item["t"], item.get("s", Pt(12)), item.get("c", WHITE), item.get("b", False))
    return tf


def box(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    rgb(shape, fill, line or LINE, Pt(1.5))
    shape.adjustments[0] = 0.08
    return shape


def chip(slide, x, y, w, h, fill, text, size=Pt(10), color=NAVY):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    rgb(shape, fill, fill)
    shape.adjustments[0] = 0.4
    add_text(shape, [{"t": text, "s": size, "c": color, "b": True}], PP_ALIGN.CENTER)
    return shape


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def header(slide, kicker, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), W, Inches(0.12))
    rgb(bar, CYAN, CYAN)
    k = slide.shapes.add_textbox(Inches(0.45), Inches(0.22), Inches(12.4), Inches(0.28))
    add_text(k, [{"t": kicker.upper(), "s": Pt(11), "c": CYAN, "b": True}])
    t = slide.shapes.add_textbox(Inches(0.45), Inches(0.44), Inches(12.4), Inches(0.42))
    add_text(t, [{"t": title, "s": Pt(26), "c": WHITE, "b": True}])
    if subtitle:
        s = slide.shapes.add_textbox(Inches(0.45), Inches(0.88), Inches(12.4), Inches(0.36))
        add_text(s, [{"t": subtitle, "s": Pt(13), "c": MUTED}])


def footer(slide, page, total=14):
    f = slide.shapes.add_textbox(Inches(0.45), Inches(7.18), Inches(10.2), Inches(0.24))
    add_text(f, [{"t": "Recon AI  ·  compose.yaml + Dockerfile  ·  evidência no repositório", "s": Pt(10), "c": MUTED}])
    n = slide.shapes.add_textbox(Inches(11.4), Inches(7.18), Inches(1.4), Inches(0.24))
    add_text(n, [{"t": f"{page} / {total}", "s": Pt(10), "c": MUTED}], PP_ALIGN.RIGHT)


def note(slide, x, y, w, h, fill, text, color=WHITE):
    shape = box(slide, x, y, w, h, fill, fill)
    add_text(shape, [{"t": text, "s": Pt(11), "c": color}])
    return shape


def new(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    return slide


def slide_cover(prs):
    s = new(prs)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), H)
    rgb(bar, CYAN, CYAN)
    k = s.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(11.5), Inches(0.35))
    add_text(k, [{"t": "ARQUITETURA DE CONTAINERS", "s": Pt(14), "c": CYAN, "b": True}])
    t = s.shapes.add_textbox(Inches(0.7), Inches(2.2), Inches(11.8), Inches(1.2))
    add_text(t, [{"t": "Recon AI — mapa Compose", "s": Pt(40), "c": WHITE, "b": True}])
    sub = s.shapes.add_textbox(Inches(0.7), Inches(3.4), Inches(11.5), Inches(0.7))
    add_text(
        sub,
        [{"t": "Oito visões: containers, rede, volumes, start-up, fluxo de dados, responsabilidades, build e configuração.", "s": Pt(16), "c": MUTED}],
    )
    for i, label in enumerate(("3 serviços", "1 rede default", "SQLite em ./results", "sem broker")):
        chip(s, Inches(0.7 + i * 2.3), Inches(4.5), Inches(2.1), Inches(0.38), CARD, label, Pt(12), WHITE)
    src = s.shapes.add_textbox(Inches(0.7), Inches(6.5), Inches(11.5), Inches(0.5))
    add_text(
        src,
        [{"t": "Fonte: compose.yaml, Dockerfile, recon_ai/{cli,ui,db,ai,control}.py  ·  Não há proxy, scheduler, gRPC nem mensageria.", "s": Pt(12), "c": MUTED}],
    )


def slide_agenda(prs):
    s = new(prs)
    header(s, "Índice", "Oito diagramas deste deck")
    items = [
        ("01", "Containers (C4)", "O que existe: recon, ui, ollama — papel, imagem, porta."),
        ("02", "Rede", "Quem chama quem, HTTP, porta, host vs rede Compose."),
        ("03", "Volumes", "Binds e volume nomeado. Dono do recon.db."),
        ("04", "Start-up", "depends_on sem healthcheck — ordem, não prontidão."),
        ("05", "Fluxo de dados", "Do spec ao SQLite, artefatos e Ollama."),
        ("06", "Responsabilidades", "Contextos de negócio, não topologia."),
        ("07", "Build", "Multi-stage Dockerfile → recon-ai:local."),
        ("08", "Configuração", "Env vars: Compose, Dockerfile, código, .env."),
    ]
    for i, (num, title, desc) in enumerate(items):
        col, row = i % 2, i // 2
        x, y = Inches(0.45 + col * 6.4), Inches(1.4 + row * 1.3)
        card = box(s, x, y, Inches(6.15), Inches(1.15), SURFACE, LINE)
        add_text(
            card,
            [
                {"t": f"{num}  {title}", "s": Pt(16), "c": CYAN, "b": True},
                {"t": desc, "s": Pt(13), "c": WHITE},
            ],
        )
    footer(s, 2)


def slide_containers(prs):
    s = new(prs)
    header(
        s,
        "1 · C4 Container",
        "O que existe no sistema",
        "Três serviços Compose. Sem proxy, worker separado, scheduler ou API FastAPI.",
    )
    host = box(s, Inches(0.4), Inches(1.4), Inches(12.5), Inches(5.5), SURFACE, LINE)
    add_text(host, [{"t": "Host do operador  ·  Docker Engine + Compose", "s": Pt(12), "c": MUTED, "b": True}])

    net = box(s, Inches(0.65), Inches(1.95), Inches(12.05), Inches(3.55), RGBColor(17, 24, 39), CYAN)
    add_text(net, [{"t": "Rede Compose padrão  ·  {projeto}_default  ·  única rede", "s": Pt(12), "c": CYAN, "b": True}])

    cards = [
        (
            0.9,
            CYAN,
            "recon",
            "WORKER / CLI",
            "Imagem recon-ai:local\nDebian 13-slim + venv Python\nuser PUID:PGID (1000:1000)\nEntrypoint: python -m recon_ai\nPorta: nenhuma no host\nSobe com compose run (efêmero)",
        ),
        (
            4.85,
            GREEN,
            "ui",
            "CONTROLE HTTP",
            "Mesma imagem recon-ai:local\nProfile: ui\nEntrypoint: python -m recon_ai.ui\n127.0.0.1:8765 → 8765\nhttp.server stdlib, HTML estático\nNão monta o código-fonte",
        ),
        (
            8.8,
            AMBER,
            "ollama",
            "MODELO LOCAL",
            "ollama/ollama@sha256:292ee7…\nSem profile (sobe sempre)\nPorta 11434 só interna\nVolume ollama_data\nModelo qwen2.5:7b\nLLAMA_ARG_CACHE_RAM=0",
        ),
    ]
    for x, color, name, role, body in cards:
        c = box(s, Inches(x), Inches(2.45), Inches(3.7), Inches(2.85), CARD, color)
        add_text(
            c,
            [
                {"t": name, "s": Pt(20), "c": color, "b": True},
                {"t": role, "s": Pt(11), "c": MUTED, "b": True},
                {"t": body, "s": Pt(12), "c": WHITE},
            ],
        )

    note(
        s,
        Inches(0.65),
        Inches(5.65),
        Inches(12.05),
        Inches(1.05),
        RGBColor(30, 41, 59),
        "Não existem: nginx/Caddy, Celery, Redis, Postgres, gRPC, webhook, cloud. A UI e o CLI compartilham a mesma imagem; só mudam entrypoint, profile e porta. A UI em :8765 precisa de docker compose build para refletir mudanças de código.",
    )
    footer(s, 3)


def slide_network(prs):
    s = new(prs)
    header(
        s,
        "2 · Rede",
        "Quem consome o quê e de onde",
        "Tudo é HTTP. Não há gRPC nem fila. Uma única rede Compose; só a UI publica porta no host.",
    )

    op = box(s, Inches(0.45), Inches(1.45), Inches(2.6), Inches(5.4), SURFACE, LINE)
    add_text(
        op,
        [
            {"t": "Operador / host", "s": Pt(14), "c": WHITE, "b": True},
            {"t": "Browser\nhttp://127.0.0.1:8765", "s": Pt(12), "c": GREEN},
            {"t": "CLI\ndocker compose run recon …", "s": Pt(12), "c": CYAN},
            {"t": "Alvo autorizado\ninternet / host.docker.internal", "s": Pt(12), "c": AMBER},
        ],
    )

    net = box(s, Inches(3.3), Inches(1.45), Inches(9.55), Inches(5.4), RGBColor(17, 24, 39), CYAN)
    add_text(net, [{"t": "rede interna  ·  {projeto}_default", "s": Pt(12), "c": CYAN, "b": True}])

    ui = box(s, Inches(3.6), Inches(2.05), Inches(3.9), Inches(2.0), CARD, GREEN)
    add_text(
        ui,
        [
            {"t": "ui :8765", "s": Pt(16), "c": GREEN, "b": True},
            {"t": "HTTP bind 0.0.0.0:8765\nPublicado só em 127.0.0.1\nGET/POST /api/*", "s": Pt(12), "c": WHITE},
        ],
    )
    rec = box(s, Inches(3.6), Inches(4.25), Inches(3.9), Inches(2.2), CARD, CYAN)
    add_text(
        rec,
        [
            {"t": "recon", "s": Pt(16), "c": CYAN, "b": True},
            {"t": "Sem porta publicada\nDNS/HTTP/HTTPS de saída\npara alvos e APIs OSINT", "s": Pt(12), "c": WHITE},
        ],
    )
    ol = box(s, Inches(8.4), Inches(2.85), Inches(4.1), Inches(2.6), CARD, AMBER)
    add_text(
        ol,
        [
            {"t": "ollama :11434", "s": Pt(16), "c": AMBER, "b": True},
            {"t": "HTTP /api/chat interno\nNÃO publicado no host\nRECON_OLLAMA_URL=\nhttp://ollama:11434", "s": Pt(12), "c": WHITE},
        ],
    )

    chip(s, Inches(7.6), Inches(2.7), Inches(0.7), Inches(0.32), GREEN, "HTTP", Pt(9), NAVY)
    chip(s, Inches(7.6), Inches(4.9), Inches(0.7), Inches(0.32), CYAN, "HTTP", Pt(9), NAVY)
    footer(s, 4)


def slide_network_table(prs):
    s = new(prs)
    header(s, "2 · Rede", "Matriz de comunicação", "Protocolo, porta, direção e se atravessa o host.")
    rows = [
        ("Origem", "Destino", "Protocolo", "Porta", "Escopo"),
        ("Browser no host", "ui", "HTTP", "127.0.0.1:8765", "Host → container (loopback)"),
        ("ui", "ollama", "HTTP", "11434", "Só rede Compose"),
        ("recon", "ollama", "HTTP", "11434", "Só rede Compose"),
        ("recon / ui", "alvo autorizado", "DNS / HTTP / HTTPS", "53, 80, 443…", "Saída para internet ou host"),
        ("recon / ui", "host.docker.internal", "HTTP", "alvo no host", "extra_hosts → host-gateway"),
        ("recon / ui", "crt.sh, gau, Bing…", "HTTPS", "443", "Saída OSINT (ferramentas)"),
        ("Operador", "recon", "CLI Compose", "—", "Não é HTTP; processo efêmero"),
    ]
    table = s.shapes.add_table(len(rows), 5, Inches(0.45), Inches(1.4), Inches(12.4), Inches(5.2)).table
    widths = [2.3, 2.5, 2.2, 2.4, 3.0]
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.name = "Calibri"
                p.font.bold = r == 0
                p.font.color.rgb = WHITE if r == 0 else WHITE
            shade = CYAN if r == 0 else (SURFACE if r % 2 else CARD)
            cell.fill.solid()
            cell.fill.fore_color.rgb = shade
    footer(s, 5)


def slide_volumes(prs):
    s = new(prs)
    header(
        s,
        "3 · Persistência",
        "Volumes, binds e o dono do SQLite",
        "SQLite é um arquivo. Dois containers montam o mesmo ./results em leitura/escrita.",
    )

    vols = [
        (CYAN, "./scope → /scope", "bind  ·  :ro", "recon + ui", "Allowlist scopes.txt. Sem este arquivo o scan recusa o alvo."),
        (GREEN, "./config → /config", "bind  ·  :ro", "recon + ui", "recon.yaml e chaves opcionais. Copiadas para state/home."),
        (AMBER, "./results → /results", "bind  ·  rw", "recon + ui", "recon.db + artefatos do run. Ponto crítico do SQLite."),
        (PURPLE, "ollama_data", "volume nomeado  ·  rw", "só ollama", "/root/.ollama — pesos do modelo (~4,7 GB)."),
    ]
    for i, (color, title, kind, who, desc) in enumerate(vols):
        y = Inches(1.35 + i * 1.15)
        card = box(s, Inches(0.45), y, Inches(12.4), Inches(1.05), SURFACE, color)
        add_text(
            card,
            [
                {"t": f"{title}    ·    {kind}    ·    {who}", "s": Pt(16), "c": color, "b": True},
                {"t": desc, "s": Pt(13), "c": WHITE},
            ],
        )
    footer(s, 6)


def slide_sqlite(prs):
    s = new(prs)
    header(s, "3 · SQLite", "Quem escreve em recon.db", "Arquivo único: /results/recon.db no bind ./results.")
    owner = box(s, Inches(0.45), Inches(1.4), Inches(6.1), Inches(5.4), SURFACE, RED)
    add_text(
        owner,
        [
            {"t": "Risco de concorrência", "s": Pt(18), "c": RED, "b": True},
            {"t": "recon e ui montam o mesmo volume rw.\nO Compose não serializa escrita.", "s": Pt(13), "c": WHITE},
            {"t": "Mitigações no código", "s": Pt(14), "c": AMBER, "b": True},
            {"t": "• UI: SCAN_LOCK no processo + recusa se já existe run running no SQLite\n• CLI: um processo por compose run\n• GET da UI usa mode=ro em view.connect_readonly\n• WAL não está ligado em db.connect\n• check_same_thread=False só para threads da UI", "s": Pt(13), "c": WHITE},
            {"t": "Dois compose run em paralelo, ou CLI + UI no mesmo db, ainda podem colidir.", "s": Pt(13), "c": AMBER},
        ],
    )
    paths = box(s, Inches(6.75), Inches(1.4), Inches(6.1), Inches(5.4), SURFACE, AMBER)
    add_text(
        paths,
        [
            {"t": "Mapa do volume ./results", "s": Pt(18), "c": AMBER, "b": True},
            {"t": "/results/recon.db\n  dono de escrita: processo do scan\n  (CLI recon ou thread da UI)", "s": Pt(13), "c": WHITE},
            {"t": "/results/recon_<alvo>_<UTC>/\n  phases.tsv, logs, JSONL, JS, HAR, corr/", "s": Pt(13), "c": WHITE},
            {"t": "/results/state/home\n  cópia gravável de /config (HOME/XDG)", "s": Pt(13), "c": WHITE},
            {"t": "ollama_data não contém o banco.\nO modelo nunca abre o .db.", "s": Pt(13), "c": GREEN},
        ],
    )
    footer(s, 7)


def slide_startup(prs):
    s = new(prs)
    header(
        s,
        "4 · Start-up",
        "Ordem de subida — sem healthcheck",
        "depends_on só espera o container iniciar, não o modelo responder.",
    )
    steps = [
        (AMBER, "1", "ollama", "Sobe primeiro\n(sem profile)\nPorta interna 11434\nSem healthcheck\nPode ainda carregar o modelo"),
        (CYAN, "2a", "recon", "depends_on: ollama\ncondition implícita:\nservice_started\ncompose run --rm"),
        (GREEN, "2b", "ui", "depends_on: ollama\nprofile ui\n127.0.0.1:8765\nTambém só service_started"),
    ]
    for i, (color, num, name, body) in enumerate(steps):
        x = Inches(0.5 + i * 4.2)
        c = box(s, x, Inches(1.5), Inches(3.95), Inches(3.3), CARD, color)
        add_text(
            c,
            [
                {"t": f"{num}  {name}", "s": Pt(20), "c": color, "b": True},
                {"t": body, "s": Pt(14), "c": WHITE},
            ],
        )
    note(
        s,
        Inches(0.5),
        Inches(5.05),
        Inches(12.3),
        Inches(1.75),
        RGBColor(69, 26, 3),
        "Não há healthcheck nem condition: service_healthy. recon/ui podem chamar http://ollama:11434 antes do qwen2.5:7b estar pronto — a fase ai falha por transporte, não por “degradação silenciosa”. UI e recon sobem em paralelo um em relação ao outro; não há depends_on entre eles. Não existe init container nem job de popular o volume.",
        WHITE,
    )
    footer(s, 8)


def slide_dataflow(prs):
    s = new(prs)
    header(
        s,
        "5 · Fluxo de dados",
        "Caminho de um scan completo",
        "Narrativa: spec autorizado → coleta → HAR/correlate → SQLite → Ollama → delta.",
    )
    steps = [
        ("1 Spec", "alvo + scopes.txt\n+ HAR opcional"),
        ("2 Coleta", "subs → dns → http\nurls → dorks → js\nnuclei"),
        ("3 HAR", "ingest captura\nnuclei sem sessão"),
        ("4 Correlate", "joins corr/…\nchecklist.json"),
        ("5 Store", "recon.db +\nartefatos /results"),
        ("6 IA", "Ollama classifica\ne extrai rotas"),
        ("7 Saída", "resumo, delta,\nPDF/Word, UI"),
    ]
    for i, (title, body) in enumerate(steps):
        x = Inches(0.35 + i * 1.85)
        c = box(s, x, Inches(1.45), Inches(1.75), Inches(1.85), CARD, CYAN if i < 5 else AMBER if i == 5 else GREEN)
        add_text(c, [{"t": title, "s": Pt(13), "c": CYAN, "b": True}, {"t": body, "s": Pt(11), "c": WHITE}])
        if i < len(steps) - 1:
            chip(s, Inches(x + 1.62), Inches(2.15), Inches(0.28), Inches(0.28), CYAN, "›", Pt(12), NAVY)

    left = box(s, Inches(0.45), Inches(3.5), Inches(6.15), Inches(3.3), SURFACE, LINE)
    add_text(
        left,
        [
            {"t": "O que atravessa a rede", "s": Pt(16), "c": CYAN, "b": True},
            {"t": "Coleta: DNS/HTTP públicos, crt.sh, gau/wayback, Bing (go-dork).", "s": Pt(13), "c": WHITE},
            {"t": "IA: prompt e JS filtrado só para ollama:11434 (rede Compose).", "s": Pt(13), "c": WHITE},
            {"t": "UI: spec JSON + HAR ≤ 32 MB no POST /api/runs.", "s": Pt(13), "c": WHITE},
            {"t": "Cloud e webhook: recusados nesta fase.", "s": Pt(13), "c": AMBER},
        ],
    )
    right = box(s, Inches(6.8), Inches(3.5), Inches(6.05), Inches(3.3), SURFACE, LINE)
    add_text(
        right,
        [
            {"t": "O que a IA vê", "s": Pt(16), "c": AMBER, "b": True},
            {"t": "Findings já ingeridos (Nuclei, har/…, corr/…) e JS em js/content com path literal.", "s": Pt(13), "c": WHITE},
            {"t": "Não lê gau.txt cru nem HTML dos alvos. Sem findings e sem path, não chama o modelo.", "s": Pt(13), "c": WHITE},
            {"t": "Schema JSON estável; o idioma do prompt segue spec.lang (pt-BR / en-US / es-ES).", "s": Pt(13), "c": WHITE},
        ],
    )
    footer(s, 9)


def slide_contexts(prs):
    s = new(prs)
    header(
        s,
        "6 · Responsabilidades",
        "Bounded contexts — visão de negócio",
        "Agrupamento conceitual. Vários contextos vivem no mesmo container recon/ui.",
    )
    groups = [
        (CYAN, "Autorização", "scopes.txt, refuse_scope, in_scope/out_scope, HAR filtrado por url_allowed. Recusa alvo fora da allowlist."),
        (GREEN, "Coleta", "13 ferramentas pinadas: enum, DNS, HTTP, crawl, arquivo, dorks, JS, Nuclei. Pipeline no processo recon/ui."),
        (AMBER, "Correlação", "Joins A∧B (HAR, JS, crt.sh, CNAME, buckets). Candidatos corr/…, não exploits."),
        (PURPLE, "Memória", "SQLite + diretório do run + delta vs último ok. Resume a partir de phases.tsv."),
        (RGBColor(244, 114, 182), "Julgamento local", "Ollama classifica findings e extrai rotas. Sem cloud. Sem webhook."),
        (RGBColor(125, 211, 252), "Controle", "CLI Typer e UI HTTP opcional. Um scan por vez na UI. Export PDF/Word."),
    ]
    for i, (color, title, body) in enumerate(groups):
        col, row = i % 3, i // 3
        x, y = Inches(0.4 + col * 4.3), Inches(1.4 + row * 2.7)
        c = box(s, x, y, Inches(4.1), Inches(2.5), SURFACE, color)
        add_text(c, [{"t": title, "s": Pt(18), "c": color, "b": True}, {"t": body, "s": Pt(13), "c": WHITE}])
    footer(s, 10)


def slide_build(prs):
    s = new(prs)
    header(
        s,
        "7 · Build",
        "Dockerfile multi-stage → imagem → container",
        "Dois estágios de build descartáveis; a imagem final não contém o compilador Go.",
    )
    stages = [
        (AMBER, "Estágio sourcebins", "golang:1.27.1-trixie@digest\nassetfinder (commit pinado)\nwaybackurls@v0.1.0\nnuclei-templates v10.4.9"),
        (GREEN, "Estágio releasetools", "debian:13-slim@digest\nscripts/fetch-tools.sh\nsubfinder amass dnsx httpx\nkatana gau nuclei go-dork"),
        (CYAN, "Imagem final", "debian:13-slim@digest\npython3-venv + requirements.txt\nCOPY bins + recon_ai\nENTRYPOINT python -m recon_ai"),
    ]
    for i, (color, title, body) in enumerate(stages):
        x = Inches(0.45 + i * 4.25)
        c = box(s, x, Inches(1.4), Inches(4.05), Inches(2.7), CARD, color)
        add_text(c, [{"t": title, "s": Pt(16), "c": color, "b": True}, {"t": body, "s": Pt(13), "c": WHITE}])
        if i < 2:
            chip(s, Inches(x + 3.85), Inches(2.5), Inches(0.4), Inches(0.32), CYAN, "→", Pt(12), NAVY)

    runtime = box(s, Inches(0.45), Inches(4.3), Inches(12.4), Inches(2.5), SURFACE, LINE)
    add_text(
        runtime,
        [
            {"t": "Runtime: a mesma tag recon-ai:local alimenta dois containers", "s": Pt(16), "c": WHITE, "b": True},
            {"t": "recon  →  entrypoint padrão (CLI).   ui  →  override python -m recon_ai.ui.   A UI não bind-mounta o source: mudar ui_page.html exige rebuild.", "s": Pt(13), "c": MUTED},
            {"t": "ARG TARGETARCH (amd64|arm64) no fetch-tools. Digests iguais a versions.lock. Cache: estágio Go e releases são independentes — mudar só Python não recompila assetfinder.", "s": Pt(13), "c": MUTED},
            {"t": "ollama não é build local: imagem puxada pelo digest. Não compartilha camadas com recon-ai:local.", "s": Pt(13), "c": MUTED},
        ],
    )
    footer(s, 11)


def slide_env_compose(prs):
    s = new(prs)
    header(s, "8 · Configuração", "Variáveis no Compose", "Origem: environment: no serviço, interpolação ${…}, ou default no YAML.")
    rows = [
        ("Serviço", "Variável", "Valor / origem", "Papel"),
        ("recon, ui", "PUID / PGID", "${PUID:-1000} (user:)", "UID dos arquivos em ./results"),
        ("recon", "RECON_ENGINE", "${RECON_ENGINE:-python}", "python (padrão) ou bash"),
        ("recon, ui", "RECON_OLLAMA_URL", "http://ollama:11434", "Cliente HTTP da IA"),
        ("recon, ui", "RECON_LOOPBACK_HOST", "host.docker.internal", "Alvo 127.0.0.1 no host"),
        ("ui", "RECON_DB", "/results/recon.db", "Caminho do SQLite"),
        ("ui", "RECON_RESULTS", "/results", "Raiz dos artefatos"),
        ("ui", "RECON_UI_HOST / PORT", "0.0.0.0 / 8765", "Bind HTTP; host publica 127.0.0.1"),
        ("ollama", "LLAMA_ARG_CACHE_RAM", "0", "Evita cache de 8 GB no runner"),
    ]
    table = s.shapes.add_table(len(rows), 4, Inches(0.4), Inches(1.35), Inches(12.5), Inches(5.35)).table
    for i, w in enumerate([1.7, 3.2, 4.0, 3.6]):
        table.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.name = "Calibri"
                p.font.bold = r == 0
                p.font.color.rgb = WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = CYAN if r == 0 else (SURFACE if r % 2 else CARD)
    footer(s, 12)


def slide_env_code(prs):
    s = new(prs)
    header(
        s,
        "8 · Configuração",
        "Dockerfile, código e arquivos — sem secrets do Compose",
        "Não há secrets: nem top-level secrets nem docker secret. Chaves vão em /config (ro) ou env do host.",
    )
    rows = [
        ("Onde", "Chave", "Default", "Notas"),
        ("Dockerfile ENV", "PYTHONPATH", "/opt/recon/app", "Pacote recon_ai"),
        ("Dockerfile ENV", "RECON_DB", "/results/recon.db", "recon herda; ui redefine"),
        ("Dockerfile ENV", "RECON_STATE_HOME", "/results/state/home", "HOME gravável"),
        ("Dockerfile ENV", "RECON_VERIFY", "verify-image.sh", "--check / start do job"),
        ("Código", "RECON_SCOPE", "/scope/scopes.txt", "Allowlist obrigatória"),
        ("Código", "RECON_CONFIG", "/config", "recon.yaml + providers"),
        ("Código", "RECON_LANG", "pt-BR se vazio", "Chrome do CLI; inválido = erro"),
        ("Código", "NUCLEI_TEMPLATES_DIR", "/opt/nuclei-templates", "Templates pinados na imagem"),
        ("Código / spec", "SHODAN_API_KEY", "vazio", "Só se dorks engine=shodan"),
        ("Arquivo", "config/recon.yaml", "threads 50…", "Bind :ro, não é env"),
        ("Arquivo", "scope/scopes.txt", "—", "Obrigatório; gitignored"),
        (".env (host)", "PUID, PGID, RECON_ENGINE", "opcional", ".gitignore ignora .env"),
    ]
    table = s.shapes.add_table(len(rows), 4, Inches(0.35), Inches(1.32), Inches(12.6), Inches(5.5)).table
    for i, w in enumerate([2.2, 3.3, 3.2, 3.9]):
        table.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.name = "Calibri"
                p.font.bold = r == 0
                p.font.color.rgb = WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = CYAN if r == 0 else (SURFACE if r % 2 else CARD)
    footer(s, 13)


def slide_absent(prs):
    s = new(prs)
    header(s, "Limites", "O que este mapa não contém — de propósito", "Útil para não procurar peças de um stack que o projeto recusou.")
    missing = [
        ("Proxy / ingress", "Nada na frente da UI. Bind só em 127.0.0.1:8765."),
        ("Scheduler", "Não há cron. Scan sob demanda (CLI ou POST)."),
        ("Fila / broker", "ADR-013: sem Celery, Redis, RabbitMQ."),
        ("Postgres", "ADR-006: um arquivo SQLite."),
        ("gRPC", "Só HTTP (UI e Ollama)."),
        ("Healthcheck", "depends_on sem condition: service_healthy."),
        ("Segunda rede", "Uma default. Sem overlay interno/externo."),
        ("Secret store", "Chaves em bind /config ou env; sem Docker secrets."),
        ("Cloud LLM", "AiError / CloudNotAvailable. Tráfego de IA fica no Compose."),
    ]
    for i, (title, body) in enumerate(missing):
        col, row = i % 3, i // 3
        x, y = Inches(0.4 + col * 4.3), Inches(1.4 + row * 1.75)
        c = box(s, x, y, Inches(4.1), Inches(1.6), SURFACE, LINE)
        add_text(c, [{"t": title, "s": Pt(15), "c": AMBER, "b": True}, {"t": body, "s": Pt(13), "c": WHITE}])
    footer(s, 14)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slide_cover(prs)
    slide_agenda(prs)
    slide_containers(prs)
    slide_network(prs)
    slide_network_table(prs)
    slide_volumes(prs)
    slide_sqlite(prs)
    slide_startup(prs)
    slide_dataflow(prs)
    slide_contexts(prs)
    slide_build(prs)
    slide_env_compose(prs)
    slide_env_code(prs)
    slide_absent(prs)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
